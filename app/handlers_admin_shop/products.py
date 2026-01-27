from aiogram import Router, F
from aiogram.types import CallbackQuery, Message, InlineKeyboardMarkup, InlineKeyboardButton
from aiogram.fsm.state import State, StatesGroup
from aiogram.fsm.context import FSMContext
from io import BytesIO
from openpyxl import load_workbook

from app.db.database import Database
from app.handlers_admin_shop.utils import get_admin_shop_ids, is_shop_admin
from app.handlers_admin_shop.start import kb_admin_main
from app.repositories.categories_repo import CategoriesRepo
from app.repositories.products_repo import ProductsRepo

router = Router()


class ProductStates(StatesGroup):
    add_category = State()
    add_product = State()
    bulk_import = State()


def kb_home() -> InlineKeyboardMarkup:
    return InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="🏠 Главная", callback_data="a:home")]
    ])


def kb_categories(cats: list[dict]) -> InlineKeyboardMarkup:
    kb = []
    for c in cats:
        status = "✅" if int(c["is_active"]) == 1 else "⛔"
        kb.append([InlineKeyboardButton(
            text=f"{status} {c['name']}",
            callback_data=f"a:pcat:{c['id']}"
        )])

    kb.append([InlineKeyboardButton(text="📥 Импорт из Excel", callback_data="a:pimport")])
    kb.append([InlineKeyboardButton(text="➕ Добавить категорию", callback_data="a:paddcat")])
    kb.append([InlineKeyboardButton(text="🏠 Главная", callback_data="a:home")])
    return InlineKeyboardMarkup(inline_keyboard=kb)


def kb_products(cat_id: int, items: list[dict]) -> InlineKeyboardMarkup:
    kb = []
    for p in items:
        status = "✅" if int(p["is_active"]) == 1 else "⛔"
        kb.append([InlineKeyboardButton(
            text=f"{status} {p['name']} — {p['price']}",
            callback_data=f"a:pprod:{cat_id}:{p['id']}"
        )])

    kb.append([InlineKeyboardButton(text="➕ Добавить товар", callback_data=f"a:paddprod:{cat_id}")])
    kb.append([
        InlineKeyboardButton(text="🔙 Категории", callback_data="a:products"),
        InlineKeyboardButton(text="🏠 Главная", callback_data="a:home"),
    ])
    return InlineKeyboardMarkup(inline_keyboard=kb)


def kb_product_card(cat_id: int, product_id: int, is_active: int) -> InlineKeyboardMarkup:
    toggle_text = "⛔ Выключить" if int(is_active) == 1 else "✅ Включить"
    return InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text=toggle_text, callback_data=f"a:ptoggle:{cat_id}:{product_id}")],
        [
            InlineKeyboardButton(text="🔙 Назад", callback_data=f"a:pcat:{cat_id}"),
            InlineKeyboardButton(text="🏠 Главная", callback_data="a:home"),
        ],
    ])


async def _get_shop_id_for_admin(db: Database, user_id: int) -> int | None:
    shop_ids = await get_admin_shop_ids(db, user_id)
    return shop_ids[0] if shop_ids else None


@router.callback_query(F.data == "a:products")
async def products_root(cq: CallbackQuery, db: Database):
    if not await is_shop_admin(db, cq.from_user.id):
        await cq.answer("Нет доступа", show_alert=True)
        return

    shop_id = await _get_shop_id_for_admin(db, cq.from_user.id)
    if not shop_id:
        await cq.message.edit_text("Нет привязанного магазина.", reply_markup=kb_home())
        await cq.answer()
        return

    async with db.conn() as conn:
        cur = await conn.execute(
            "SELECT id, name, sort, is_active FROM categories WHERE shop_id=? ORDER BY sort, id",
            (shop_id,),
        )
        cats = [dict(r) for r in await cur.fetchall()]

    if not cats:
        await cq.message.edit_text(
            "🧺 Продукты\n\nКатегорий пока нет.\nНажми «➕ Добавить категорию».",
            reply_markup=InlineKeyboardMarkup(inline_keyboard=[
                [InlineKeyboardButton(text="📥 Импорт из Excel", callback_data="a:pimport")],
                [InlineKeyboardButton(text="➕ Добавить категорию", callback_data="a:paddcat")],
                [InlineKeyboardButton(text="🏠 Главная", callback_data="a:home")],
            ])
        )
        await cq.answer()
        return

    await cq.message.edit_text("🧺 Категории:", reply_markup=kb_categories(cats))
    await cq.answer()


@router.callback_query(F.data == "a:paddcat")
async def add_category_prompt(cq: CallbackQuery, state: FSMContext, db: Database):
    if not await is_shop_admin(db, cq.from_user.id):
        await cq.answer("Нет доступа", show_alert=True)
        return

    await state.set_state(ProductStates.add_category)
    await cq.message.edit_text("Введите название категории (например: Овощи):", reply_markup=kb_home())
    await cq.answer()


@router.message(ProductStates.add_category)
async def add_category_save(message: Message, state: FSMContext, db: Database):
    if not await is_shop_admin(db, message.from_user.id):
        await message.answer("Нет доступа.")
        return

    shop_id = await _get_shop_id_for_admin(db, message.from_user.id)
    if not shop_id:
        await message.answer("Нет привязанного магазина.", reply_markup=kb_admin_main())
        await state.clear()
        return

    name = (message.text or "").strip()
    if len(name) < 2:
        await message.answer("Слишком коротко. Введите название категории ещё раз:")
        return

    async with db.conn() as conn:
        await conn.execute(
            "INSERT INTO categories(shop_id, name, sort, is_active) VALUES(?,?,0,1)",
            (shop_id, name),
        )
        await conn.commit()

    await state.clear()
    await message.answer("Категория добавлена ✅", reply_markup=kb_admin_main())


@router.callback_query(F.data.startswith("a:pcat:"))
async def open_category(cq: CallbackQuery, db: Database):
    if not await is_shop_admin(db, cq.from_user.id):
        await cq.answer("Нет доступа", show_alert=True)
        return

    shop_id = await _get_shop_id_for_admin(db, cq.from_user.id)
    if not shop_id:
        await cq.message.edit_text("Нет привязанного магазина.", reply_markup=kb_home())
        await cq.answer()
        return

    cat_id = int(cq.data.split(":")[2])

    repo = ProductsRepo(db)
    items = await repo.list_by_category_any(shop_id=shop_id, category_id=cat_id)

    title = f"🧺 Товары в категории #{cat_id}:"
    await cq.message.edit_text(title, reply_markup=kb_products(cat_id, items))
    await cq.answer()


@router.callback_query(F.data.startswith("a:paddprod:"))
async def add_product_prompt(cq: CallbackQuery, state: FSMContext, db: Database):
    if not await is_shop_admin(db, cq.from_user.id):
        await cq.answer("Нет доступа", show_alert=True)
        return

    cat_id = int(cq.data.split(":")[2])
    await state.set_state(ProductStates.add_product)
    await state.update_data(category_id=cat_id)
    await cq.message.edit_text(
        "Введите товар в формате:\n"
        "Название; Цена\n\n"
        "Пример:\n"
        "Молоко 2.5%; 12.5",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=[
            [InlineKeyboardButton(text="🔙 Назад", callback_data=f"a:pcat:{cat_id}")],
            [InlineKeyboardButton(text="🏠 Главная", callback_data="a:home")],
        ])
    )
    await cq.answer()


@router.message(ProductStates.add_product)
async def add_product_save(message: Message, state: FSMContext, db: Database):
    if not await is_shop_admin(db, message.from_user.id):
        await message.answer("Нет доступа.")
        return

    shop_id = await _get_shop_id_for_admin(db, message.from_user.id)
    if not shop_id:
        await message.answer("Нет привязанного магазина.", reply_markup=kb_admin_main())
        await state.clear()
        return

    data = await state.get_data()
    cat_id = int(data["category_id"])

    text = (message.text or "").strip()
    if ";" not in text:
        await message.answer("Неверный формат. Нужно: Название; Цена\nПример: Хлеб; 7")
        return

    name, price_str = [x.strip() for x in text.split(";", 1)]
    try:
        price = float(price_str.replace(",", "."))
    except ValueError:
        await message.answer("Цена должна быть числом. Пример: 12.5")
        return

    repo = ProductsRepo(db)
    await repo.create(shop_id=shop_id, category_id=cat_id, name=name, price=price)

    await state.clear()
    await message.answer("Товар добавлен ✅", reply_markup=kb_admin_main())


def _normalize_header(value: str | None) -> str:
    if value is None:
        return ""
    return str(value).strip().lower()


@router.callback_query(F.data == "a:pimport")
async def import_products_prompt(cq: CallbackQuery, state: FSMContext, db: Database):
    if not await is_shop_admin(db, cq.from_user.id):
        await cq.answer("Нет доступа", show_alert=True)
        return

    await state.set_state(ProductStates.bulk_import)
    await cq.message.edit_text(
        "📥 Импорт товаров из Excel\n\n"
        "Пришлите файл .xlsx со столбцами:\n"
        "Категория | Название | Цена | Local_name/Err_name (опционально) | Описание (опционально)\n\n"
        "Категория может быть названием или ID.\n"
        "Пример заголовков: category, name, price, local_name, description.",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=[
            [InlineKeyboardButton(text="🏠 Главная", callback_data="a:home")],
        ]),
    )
    await cq.answer()


@router.message(ProductStates.bulk_import, F.document)
async def import_products_from_excel(message: Message, state: FSMContext, db: Database):
    if not await is_shop_admin(db, message.from_user.id):
        await message.answer("Нет доступа.")
        return

    shop_id = await _get_shop_id_for_admin(db, message.from_user.id)
    if not shop_id:
        await message.answer("Нет привязанного магазина.", reply_markup=kb_admin_main())
        await state.clear()
        return

    document = message.document
    if not document or not document.file_name.lower().endswith(".xlsx"):
        await message.answer("Нужен файл .xlsx. Попробуйте ещё раз.")
        return

    file = await message.bot.get_file(document.file_id)
    buffer = BytesIO()
    await message.bot.download_file(file.file_path, destination=buffer)
    buffer.seek(0)

    workbook = load_workbook(buffer, data_only=True)
    sheet = workbook.active

    header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
    if not header_row:
        await message.answer("Файл пустой или без заголовков.")
        return

    headers = [_normalize_header(value) for value in header_row]

    def find_col(options: set[str]) -> int | None:
        for idx, header in enumerate(headers):
            if header in options:
                return idx
        return None

    category_idx = find_col({"category", "категория", "категории"})
    name_idx = find_col({"name", "название", "товар"})
    price_idx = find_col({"price", "цена", "стоимость"})
    local_name_idx = find_col({"local_name", "local", "err_name", "локальное название", "локальное имя"})
    description_idx = find_col({"description", "описание"})

    if category_idx is None or name_idx is None or price_idx is None:
        await message.answer(
            "Не найден один из обязательных столбцов: категория, название, цена."
        )
        return

    async with db.conn() as conn:
        cur = await conn.execute(
            "SELECT id, name FROM categories WHERE shop_id=?",
            (shop_id,),
        )
        rows = await cur.fetchall()

    categories_by_name = {row["name"].strip().lower(): row["id"] for row in rows}
    categories_by_id = {int(row["id"]) for row in rows}
    categories_repo = CategoriesRepo(db)
    products_repo = ProductsRepo(db)

    created = 0
    created_categories = 0
    errors: list[str] = []

    for row_index, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if not row or all(cell is None or str(cell).strip() == "" for cell in row):
            continue

        category_cell = row[category_idx] if category_idx < len(row) else None
        name_cell = row[name_idx] if name_idx < len(row) else None
        price_cell = row[price_idx] if price_idx < len(row) else None
        local_name_cell = row[local_name_idx] if local_name_idx is not None and local_name_idx < len(row) else None
        description_cell = row[description_idx] if description_idx is not None and description_idx < len(row) else None

        if category_cell is None or str(category_cell).strip() == "":
            errors.append(f"Строка {row_index}: нет категории")
            continue

        category_id: int | None = None
        category_raw = str(category_cell).strip()
        if category_raw.isdigit():
            candidate_id = int(category_raw)
            if candidate_id in categories_by_id:
                category_id = candidate_id
            else:
                errors.append(f"Строка {row_index}: категория ID {candidate_id} не найдена")
                continue
        else:
            normalized = category_raw.lower()
            category_id = categories_by_name.get(normalized)
            if category_id is None:
                category_id = await categories_repo.create(shop_id=shop_id, name=category_raw)
                categories_by_name[normalized] = category_id
                categories_by_id.add(category_id)
                created_categories += 1

        if name_cell is None or str(name_cell).strip() == "":
            errors.append(f"Строка {row_index}: нет названия")
            continue

        name = str(name_cell).strip()
        local_name = str(local_name_cell).strip() if local_name_cell is not None else None
        description = str(description_cell).strip() if description_cell is not None else None

        if price_cell is None or str(price_cell).strip() == "":
            errors.append(f"Строка {row_index}: нет цены")
            continue

        try:
            price = float(str(price_cell).replace(",", "."))
        except ValueError:
            errors.append(f"Строка {row_index}: цена не число")
            continue

        await products_repo.create(
            shop_id=shop_id,
            category_id=category_id,
            name=name,
            price=price,
            description=description,
            local_name=local_name if local_name else None,
        )
        created += 1

    await state.clear()
    error_text = ""
    if errors:
        preview = "\n".join(errors[:10])
        suffix = "\n... и ещё ошибки." if len(errors) > 10 else ""
        error_text = f"\n\nОшибки:\n{preview}{suffix}"

    await message.answer(
        f"Импорт завершён ✅\n"
        f"Добавлено товаров: {created}\n"
        f"Создано категорий: {created_categories}"
        f"{error_text}",
        reply_markup=kb_admin_main(),
    )


@router.message(ProductStates.bulk_import)
async def import_products_wrong_message(message: Message):
    await message.answer("Пришлите файл .xlsx для импорта товаров.")


@router.callback_query(F.data.startswith("a:pprod:"))
async def product_card(cq: CallbackQuery, db: Database):
    if not await is_shop_admin(db, cq.from_user.id):
        await cq.answer("Нет доступа", show_alert=True)
        return

    _, _, cat_id_s, prod_id_s = cq.data.split(":")
    cat_id = int(cat_id_s)
    product_id = int(prod_id_s)

    repo = ProductsRepo(db)
    p = await repo.get(product_id)
    if not p:
        await cq.message.edit_text("Товар не найден.", reply_markup=kb_home())
        await cq.answer()
        return

    text = (
        f"📦 {p['name']}\n"
        f"Цена: {p['price']}\n"
        f"Статус: {'✅ Активен' if int(p['is_active'])==1 else '⛔ Выключен'}\n"
    )
    await cq.message.edit_text(text, reply_markup=kb_product_card(cat_id, product_id, p["is_active"]))
    await cq.answer()


@router.callback_query(F.data.startswith("a:ptoggle:"))
async def toggle_product(cq: CallbackQuery, db: Database):
    if not await is_shop_admin(db, cq.from_user.id):
        await cq.answer("Нет доступа", show_alert=True)
        return

    _, _, cat_id_s, prod_id_s = cq.data.split(":")
    cat_id = int(cat_id_s)
    product_id = int(prod_id_s)

    shop_id = await _get_shop_id_for_admin(db, cq.from_user.id)
    if not shop_id:
        await cq.answer("Нет магазина", show_alert=True)
        return

    repo = ProductsRepo(db)
    new_state = await repo.toggle_active(shop_id=shop_id, product_id=product_id)
    if new_state is None:
        await cq.answer("Товар не найден", show_alert=True)
        return

    await cq.answer("Готово ✅")
    # перерисуем карточку
    p = await repo.get(product_id)
    text = (
        f"📦 {p['name']}\n"
        f"Цена: {p['price']}\n"
        f"Статус: {'✅ Активен' if int(p['is_active'])==1 else '⛔ Выключен'}\n"
    )
    await cq.message.edit_text(text, reply_markup=kb_product_card(cat_id, product_id, p["is_active"]))
